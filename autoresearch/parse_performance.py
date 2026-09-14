# -*- coding: utf-8 -*-
"""
从 performance_metrics.xlsx 解析关键指标（阶段 1）。

Excel 由 playwright 保存，每个 sheet 对应一项风险指标，列一般为：
日期 | 1个月 | 3个月 | 6个月 | 12个月
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PERIOD_HEADER_MAP = {
    "1个月": "1month",
    "3个月": "3month",
    "6个月": "6month",
    "12个月": "12month",
}

SHEET_STRATEGY_RETURN = "策略收益"
SHEET_MAX_DRAWDOWN = "最大回撤"
SHEET_SHARPE = "夏普比率"
SHEET_ALPHA = "阿尔法"

PCT_RE = re.compile(r"-?\d+(?:\.\d+)?")


def _ratio_to_percent(value: float | None) -> float | None:
    """聚宽部分指标为小数（0.0705 表示 7.05%），统一为百分数便于与 goal 比较。"""
    if value is None:
        return None
    if abs(value) <= 1.5:
        return round(value * 100, 4)
    return value


def _parse_percent(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("％", "%")
    if not text or text in ("--", "-", "N/A", "nan"):
        return None
    if text.endswith("%"):
        try:
            return float(text[:-1])
        except ValueError:
            pass
    m = PCT_RE.search(text)
    if not m:
        return None
    try:
        return float(m.group(0))
    except ValueError:
        return None


def _read_sheet_rows(ws) -> tuple[list[str], list[tuple[Any, ...]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    data = [r for r in rows[1:] if any(c is not None and str(c).strip() for c in r)]
    return headers, data


def _period_indices(headers: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, h in enumerate(headers):
        key = PERIOD_HEADER_MAP.get(h)
        if key:
            out[key] = idx
    return out


def _pick_primary_period(by_period: dict[str, float | None]) -> tuple[str | None, float | None]:
    """回测区间多为半年左右，优先 6 个月，其次 12/3/1。"""
    for key in ("6month", "12month", "3month", "1month"):
        val = by_period.get(key)
        if val is not None:
            return key, val
    return None, None


def _extract_period_metrics(ws) -> dict[str, Any]:
    headers, data = _read_sheet_rows(ws)
    period_idx = _period_indices(headers)
    if not data or not period_idx:
        return {"by_period": {}, "primary_period": None, "primary_value": None}

    by_period: dict[str, float | None] = {k: None for k in period_idx}
    # 最后一行常为当前回测区间汇总
    last = data[-1]
    for key, idx in period_idx.items():
        if idx < len(last):
            by_period[key] = _parse_percent(last[idx])

    primary_period, primary_value = _pick_primary_period(by_period)

    # 最大回撤：取各周期列中最「深」的一档（绝对值最大）
    worst: float | None = None
    for row in data:
        for key, idx in period_idx.items():
            if idx >= len(row):
                continue
            v = _parse_percent(row[idx])
            if v is None:
                continue
            depth = abs(v)
            if worst is None or depth > worst:
                worst = depth

    return {
        "by_period": by_period,
        "primary_period": primary_period,
        "primary_value": primary_value,
        "worst_abs": worst,
    }


def parse_performance_metrics(xlsx_path: Path | str) -> dict[str, Any]:
    """
    解析 performance_metrics.xlsx，返回统一 metrics 字典。

    annual_return_pct：来自「策略收益」主周期列
    max_drawdown_pct：来自「最大回撤」各周期绝对值最大（若无则用主周期）
    """
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(f"未找到性能分析文件：{path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        metrics: dict[str, Any] = {
            "annual_return_pct": None,
            "max_drawdown_pct": None,
            "sharpe": None,
            "alpha": None,
            "raw": {},
        }

        if SHEET_STRATEGY_RETURN in wb.sheetnames:
            sr = _extract_period_metrics(wb[SHEET_STRATEGY_RETURN])
            metrics["raw"][SHEET_STRATEGY_RETURN] = sr
            metrics["annual_return_pct"] = _ratio_to_percent(sr.get("primary_value"))

        if SHEET_MAX_DRAWDOWN in wb.sheetnames:
            dd = _extract_period_metrics(wb[SHEET_MAX_DRAWDOWN])
            metrics["raw"][SHEET_MAX_DRAWDOWN] = dd
            dd_val = dd.get("worst_abs")
            if dd_val is None:
                pv = dd.get("primary_value")
                dd_val = abs(pv) if pv is not None else None
            metrics["max_drawdown_pct"] = _ratio_to_percent(dd_val)

        if SHEET_SHARPE in wb.sheetnames:
            sh = _extract_period_metrics(wb[SHEET_SHARPE])
            metrics["raw"][SHEET_SHARPE] = sh
            metrics["sharpe"] = sh.get("primary_value")

        if SHEET_ALPHA in wb.sheetnames:
            al = _extract_period_metrics(wb[SHEET_ALPHA])
            metrics["raw"][SHEET_ALPHA] = al
            metrics["alpha"] = al.get("primary_value")

        return metrics
    finally:
        wb.close()


def parse_result_dir(result_dir: Path | str) -> dict[str, Any]:
    """从一次回测结果目录解析 metrics（要求含 performance_metrics.xlsx）。"""
    d = Path(result_dir)
    xlsx = d / "performance_metrics.xlsx"
    return parse_performance_metrics(xlsx)
